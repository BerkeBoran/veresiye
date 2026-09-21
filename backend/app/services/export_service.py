from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font


def build_statement_excel(customer, sales, payments, totals) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ekstre"
    bold = Font(bold=True)

    ws.append([f"{customer.name} {customer.surname or ''} — Hesap Ekstresi".strip()])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])

    ws.append(["Tarih", "Belge No", "Açıklama", "Ürünler", "Tutar", "Ödeme", "Borç"])
    for cell in ws[ws.max_row]:
        cell.font = bold

    movements = [("sale", s.created_at, s) for s in sales] + \
                [("payment", p.created_at, p) for p in payments]
    movements.sort(key=lambda m: m[1])

    balance = 0.0
    for kind, date, obj in movements:
        if kind == "sale":
            balance += float(obj.total)
            products = ", ".join(i.product_name for i in obj.items)
            ws.append([
                date.strftime("%d.%m.%Y"),
                obj.document_no or "",
                obj.note or "",
                products,
                float(obj.total),
                "",
                balance,
            ])
        else:
            balance -= float(obj.amount)
            ws.append([
                date.strftime("%d.%m.%Y"),
                "",
                obj.note or "",
                "",
                "",
                float(obj.amount),
                balance,
            ])

    ws.append([])

    ws.append(["Toplam Satış", "", "", "", float(totals["total_sales"])])
    ws.append(["Toplam Tahsilat", "", "", "", "", float(totals["total_payments"])])
    ws.append(["Kalan Borç", "", "", "", "", "", float(totals["balance"])])
    for r in range(ws.max_row - 2, ws.max_row + 1):
        ws[f"A{r}"].font = bold

    for col, width in zip("ABCDEFG", [12, 12, 24, 26, 12, 12, 12]):
        ws.column_dimensions[col].width = width

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()